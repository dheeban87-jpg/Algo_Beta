"""
PAPER_TRADE_LOGGER.PY — Unified Paper Trade Excel Logger
=========================================================
Central logger that all phases call to record paper trades
into the master Excel workbook (Paper_Trade_Log.xlsx).

Usage (from any phase):
    from paper_trade_logger import log_paper_entry, log_paper_exit

    # When a paper trade is entered:
    trade_id = log_paper_entry(
        phase="PH5",
        symbol="HDFC",
        direction="BUY",
        entry_price=1520.00,
        quantity=10,
        stop_loss=1494.40,
        target=1565.60,
        strategy="Gap Down Recovery",
        score=82,
        regime="RANGE_BOUND",
        notes="FinBERT BULLISH, ADX=18"
    )

    # When the trade closes:
    log_paper_exit(
        trade_id=trade_id,
        exit_price=1558.00,
        exit_reason="TARGET HIT"
    )

Brokerage model (v2.0.0 — NSE MIS, Zerodha-style):
    Brokerage  : 0.03% per leg, capped at ₹20/order
    STT        : 0.025% on sell-side turnover
    Exchange   : 0.00325% per leg (NSE transaction charge)
    SEBI       : 0.0001% per leg
    Stamp duty : 0.003% on buy-side turnover
    GST        : 18% on (brokerage + exchange + SEBI)

    For a typical ₹15,000 position, total round-trip cost ≈ ₹16
    Net P&L = Gross P&L − total brokerage
    Status (WIN/LOSS) is determined from Net P&L.

Author: Dheebanraj + Claude
Version: 2.0.0 — Brokerage added: 2026-04-22
"""

import os
import json
import logging
import threading
from datetime import datetime
from typing import Optional

logger = logging.getLogger("PaperTradeLogger")

# ── Path to the master Excel workbook (matches config.PAPER_TRADE_LOG_XLSX) ──
try:
    from config import Config
    XLSX_PATH = Config.PAPER_TRADE_LOG_XLSX
except Exception:
    XLSX_PATH = "data/Paper_Trade_Log.xlsx"

# ── Companion JSON for fast in-process lookups ──
JSON_PATH = XLSX_PATH.replace(".xlsx", "_index.json")

_lock = threading.Lock()   # thread-safe writes across phases


# ════════════════════════════════════════════════════════════════════════════
# BROKERAGE CALCULATOR (v2.0.0)
# ════════════════════════════════════════════════════════════════════════════

def _calculate_brokerage(direction: str, entry_price: float,
                         exit_price: float, quantity: int) -> dict:
    """
    Calculate realistic NSE MIS intraday brokerage (Zerodha-style).

    For BUY direction (gap-down trade):  entry = BUY leg,  exit = SELL leg
    For SELL direction (gap-up trade):   entry = SELL leg, exit = BUY  leg

    Returns a dict with per-charge breakdown and rounded total.
    """
    entry_turnover = entry_price * quantity
    exit_turnover  = exit_price  * quantity

    if direction.upper() == "BUY":
        buy_turnover  = entry_turnover
        sell_turnover = exit_turnover
    else:
        sell_turnover = entry_turnover
        buy_turnover  = exit_turnover

    # Brokerage: 0.03% per leg, capped at ₹20/order
    brok_entry = min(entry_turnover * 0.0003, 20.0)
    brok_exit  = min(exit_turnover  * 0.0003, 20.0)
    brok_total = brok_entry + brok_exit

    # STT: 0.025% on sell-side turnover (MIS intraday)
    stt = sell_turnover * 0.00025

    # NSE exchange transaction charge: 0.00325% per leg
    exch = (entry_turnover + exit_turnover) * 0.0000325

    # SEBI charges: 0.0001% per leg
    sebi = (entry_turnover + exit_turnover) * 0.000001

    # Stamp duty: 0.003% on buy-side turnover
    stamp = buy_turnover * 0.00003

    # GST: 18% on brokerage + exchange + SEBI
    gst = (brok_total + exch + sebi) * 0.18

    total = brok_total + stt + exch + sebi + stamp + gst

    return {
        "brokerage": round(brok_total, 2),
        "stt":       round(stt, 2),
        "exchange":  round(exch, 2),
        "sebi":      round(sebi, 2),
        "stamp":     round(stamp, 2),
        "gst":       round(gst, 2),
        "total":     round(total, 2),
    }


# ════════════════════════════════════════════════════════════════════════════
# PUBLIC API
# ════════════════════════════════════════════════════════════════════════════

def log_paper_entry(
    phase: str,
    symbol: str,
    direction: str,          # "BUY" or "SELL"
    entry_price: float,
    quantity: int,
    stop_loss: float,
    target: float,
    strategy: str = "",
    score: float = 0,
    regime: str = "",
    notes: str = ""
) -> str:
    """
    Record a new paper trade entry.
    Returns a unique trade_id (e.g. "PH5-HDFC-20260420-001").
    Call log_paper_exit() when the position closes.
    """
    now = datetime.now()
    date_str = now.strftime("%Y-%m-%d")
    time_str = now.strftime("%H:%M:%S")
    trade_id = _generate_trade_id(phase, symbol, now)

    risk   = abs(entry_price - stop_loss) * quantity
    reward = abs(target - entry_price) * quantity
    rr_ratio = round(reward / risk, 2) if risk > 0 else 0

    row = {
        "trade_id":     trade_id,
        "date":         date_str,
        "entry_time":   time_str,
        "exit_time":    "",
        "phase":        phase,
        "symbol":       symbol,
        "direction":    direction.upper(),
        "entry_price":  round(entry_price, 2),
        "exit_price":   "",
        "stop_loss":    round(stop_loss, 2),
        "target":       round(target, 2),
        "quantity":     quantity,
        "pnl_rs":       "",          # gross P&L — filled at exit
        "pnl_pct":      "",          # gross P&L % — filled at exit
        "brokerage_rs": "",          # total brokerage — filled at exit
        "net_pnl_rs":   "",          # net P&L after brokerage — filled at exit
        "net_pnl_pct":  "",          # net P&L % — filled at exit
        "risk_rs":      round(risk, 2),
        "reward_rs":    round(reward, 2),
        "rr_ratio":     rr_ratio,
        "strategy":     strategy,
        "score":        score,
        "regime":       regime,
        "exit_reason":  "",
        "status":       "OPEN",
        "notes":        notes,
    }

    with _lock:
        _append_row(row)

    logger.info(
        f"[PAPER] ENTRY logged → {trade_id}  "
        f"{symbol} {direction} @ {entry_price}  SL={stop_loss}  T={target}"
    )
    return trade_id


def log_paper_exit(
    trade_id: str,
    exit_price: float,
    exit_reason: str = ""
) -> dict:
    """
    Update an existing paper trade with exit details.
    Calculates gross P&L, brokerage breakdown, and net P&L.
    Status (WIN/LOSS/BREAKEVEN) is determined from net P&L.
    Returns the completed trade row dict.
    """
    now = datetime.now()
    exit_time = now.strftime("%H:%M:%S")

    with _lock:
        rows = _load_index()
        if trade_id not in rows:
            logger.warning(f"[PAPER] trade_id not found in index: {trade_id}")
            return {}

        row          = rows[trade_id]
        entry_price  = row["entry_price"]
        quantity     = row["quantity"]
        direction    = row["direction"]

        # ── Gross P&L ──
        if direction == "BUY":
            pnl_rs = (exit_price - entry_price) * quantity
        else:
            pnl_rs = (entry_price - exit_price) * quantity
        pnl_pct = (pnl_rs / (entry_price * quantity)) * 100 if entry_price > 0 else 0

        # ── Brokerage breakdown ──
        brok = _calculate_brokerage(direction, entry_price, exit_price, quantity)
        brokerage_rs = brok["total"]

        # ── Net P&L (after brokerage) ──
        net_pnl_rs  = pnl_rs - brokerage_rs
        net_pnl_pct = (net_pnl_rs / (entry_price * quantity)) * 100 if entry_price > 0 else 0

        row["exit_price"]   = round(exit_price, 2)
        row["exit_time"]    = exit_time
        row["pnl_rs"]       = round(pnl_rs, 2)
        row["pnl_pct"]      = round(pnl_pct, 4)
        row["brokerage_rs"] = round(brokerage_rs, 2)
        row["net_pnl_rs"]   = round(net_pnl_rs, 2)
        row["net_pnl_pct"]  = round(net_pnl_pct, 4)
        row["exit_reason"]  = exit_reason
        # Status is based on NET P&L (after brokerage)
        row["status"] = (
            "WIN"       if net_pnl_rs > 0  else
            "LOSS"      if net_pnl_rs < 0  else
            "BREAKEVEN"
        )

        rows[trade_id] = row
        _save_index(rows)
        _rebuild_xlsx(rows)

    logger.info(
        f"[PAPER] EXIT logged → {trade_id}  exit={exit_price}  "
        f"Gross=₹{pnl_rs:.2f}  Brok=₹{brokerage_rs:.2f}  "
        f"Net=₹{net_pnl_rs:.2f} ({net_pnl_pct:.2f}%)  [{exit_reason}]"
    )
    return row


def get_open_positions(phase: Optional[str] = None) -> list:
    """Return all currently open paper positions (optionally filtered by phase)."""
    with _lock:
        rows = _load_index()
    open_trades = [r for r in rows.values() if r["status"] == "OPEN"]
    if phase:
        open_trades = [r for r in open_trades if r["phase"] == phase]
    return open_trades


def get_daily_summary(date_str: Optional[str] = None) -> dict:
    """Return P&L summary for a given date (default: today).
    Returns gross P&L, brokerage paid, and net P&L separately.
    """
    if not date_str:
        date_str = datetime.now().strftime("%Y-%m-%d")
    with _lock:
        rows = _load_index()
    day_trades = [r for r in rows.values()
                  if r["date"] == date_str and r["status"] != "OPEN"]
    wins   = [r for r in day_trades if r["status"] == "WIN"]
    losses = [r for r in day_trades if r["status"] == "LOSS"]

    gross_pnl = sum(r["pnl_rs"]
                    for r in day_trades if isinstance(r.get("pnl_rs"), (int, float)))
    brokerage = sum(r["brokerage_rs"]
                    for r in day_trades if isinstance(r.get("brokerage_rs"), (int, float)))
    net_pnl   = sum(r["net_pnl_rs"]
                    for r in day_trades if isinstance(r.get("net_pnl_rs"), (int, float)))
    # Fallback: if net_pnl_rs not yet in old records, use gross
    if net_pnl == 0 and gross_pnl != 0:
        net_pnl = gross_pnl

    return {
        "date":       date_str,
        "trades":     len(day_trades),
        "wins":       len(wins),
        "losses":     len(losses),
        "win_rate":   round(len(wins) / len(day_trades) * 100, 1) if day_trades else 0,
        "gross_pnl":  round(gross_pnl, 2),
        "brokerage":  round(brokerage, 2),
        "total_pnl":  round(net_pnl, 2),   # "total_pnl" kept for compat; now = net
    }


# ════════════════════════════════════════════════════════════════════════════
# INTERNAL HELPERS
# ════════════════════════════════════════════════════════════════════════════

def _generate_trade_id(phase: str, symbol: str, dt: datetime) -> str:
    date_part = dt.strftime("%Y%m%d")
    rows = _load_index()
    serial = sum(1 for r in rows.values()
                 if r["phase"] == phase and r["date"] == dt.strftime("%Y-%m-%d")) + 1
    return f"{phase}-{symbol}-{date_part}-{serial:03d}"


def _load_index() -> dict:
    """Load the JSON index {trade_id: row_dict}."""
    os.makedirs(os.path.dirname(JSON_PATH), exist_ok=True)
    if not os.path.exists(JSON_PATH):
        return {}
    try:
        with open(JSON_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def _save_index(rows: dict):
    os.makedirs(os.path.dirname(JSON_PATH), exist_ok=True)
    with open(JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(rows, f, indent=2, default=str)


def _append_row(row: dict):
    """Append a new entry to index + rebuild Excel."""
    rows = _load_index()
    rows[row["trade_id"]] = row
    _save_index(rows)
    _rebuild_xlsx(rows)


def _effective_pnl(trade: dict) -> float:
    """Return net_pnl_rs if available, else gross pnl_rs, else 0."""
    net = trade.get("net_pnl_rs")
    if isinstance(net, (int, float)):
        return net
    gross = trade.get("pnl_rs")
    if isinstance(gross, (int, float)):
        return gross
    return 0.0


def _rebuild_xlsx(rows: dict):
    """
    Rebuild the full Excel workbook from the index.
    Called after every write so the file stays current.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.warning("[PAPER] openpyxl not installed — Excel log not updated. "
                       "Run: pip install openpyxl")
        return

    os.makedirs(os.path.dirname(XLSX_PATH), exist_ok=True)

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Trade Log"
    _build_trade_log_sheet(ws, rows)

    ws2 = wb.create_sheet("Daily Summary")
    _build_daily_summary_sheet(ws2, rows)

    ws3 = wb.create_sheet("Phase Performance")
    _build_phase_performance_sheet(ws3, rows)

    ws4 = wb.create_sheet("Month Overview")
    _build_month_overview_sheet(ws4, rows)

    wb.save(XLSX_PATH)
    logger.debug(f"[PAPER] Excel log saved → {XLSX_PATH}")


# ════════════════════════════════════════════════════════════════════════════
# SHEET BUILDERS
# ════════════════════════════════════════════════════════════════════════════

def _build_trade_log_sheet(ws, rows: dict):
    """Write all trade rows to the Trade Log sheet (v2.0.0 — with brokerage)."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    HEADERS = [
        "Trade ID", "Date", "Entry Time", "Exit Time",
        "Phase", "Symbol", "Direction",
        "Entry ₹", "Exit ₹", "Stop Loss ₹", "Target ₹", "Qty",
        "Gross P&L (₹)", "Gross P&L (%)",
        "Brokerage (₹)",
        "Net P&L (₹)", "Net P&L (%)",
        "Risk (₹)", "Reward (₹)", "R:R",
        "Strategy", "Score", "Regime", "Exit Reason", "Status", "Notes"
    ]
    COL_MAP = [
        "trade_id", "date", "entry_time", "exit_time",
        "phase", "symbol", "direction",
        "entry_price", "exit_price", "stop_loss", "target", "quantity",
        "pnl_rs", "pnl_pct",
        "brokerage_rs",
        "net_pnl_rs", "net_pnl_pct",
        "risk_rs", "reward_rs", "rr_ratio",
        "strategy", "score", "regime", "exit_reason", "status", "notes"
    ]

    hdr_fill = PatternFill("solid", start_color="1F4E79")
    hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    thin     = Side(style="thin", color="BFBFBF")
    border   = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border    = border

    ws.row_dimensions[1].height = 30

    WIN_FILL  = PatternFill("solid", start_color="E2EFDA")   # light green
    LOSS_FILL = PatternFill("solid", start_color="FCE4D6")   # light red
    OPEN_FILL = PatternFill("solid", start_color="FFF2CC")   # light yellow
    BASE_FONT = Font(name="Arial", size=9)

    sorted_rows = sorted(rows.values(),
                         key=lambda r: (r.get("date", ""), r.get("entry_time", "")))

    for row_idx, trade in enumerate(sorted_rows, start=2):
        status = trade.get("status", "OPEN")
        if status == "WIN":
            row_fill = WIN_FILL
        elif status == "LOSS":
            row_fill = LOSS_FILL
        else:
            row_fill = OPEN_FILL

        for col_idx, key in enumerate(COL_MAP, start=1):
            val = trade.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font      = BASE_FONT
            cell.fill      = row_fill
            cell.border    = border
            cell.alignment = Alignment(vertical="center")

            # ── Numeric formatting ──
            if key == "pnl_rs" and isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
                cell.font = Font(name="Arial", size=9,
                                 color="595959")          # grey — secondary metric

            elif key == "pnl_pct" and isinstance(val, (int, float)):
                cell.number_format = '0.00"%"'
                cell.font = Font(name="Arial", size=9, color="595959")

            elif key == "brokerage_rs" and isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
                cell.font = Font(name="Arial", size=9, color="9C5700")  # orange

            elif key == "net_pnl_rs" and isinstance(val, (int, float)):
                # Primary P&L — bold, coloured
                cell.number_format = '₹#,##0.00;[Red](₹#,##0.00)'
                cell.font = Font(name="Arial", size=9, bold=True,
                                 color="375623" if val >= 0 else "9C0006")

            elif key == "net_pnl_pct" and isinstance(val, (int, float)):
                cell.number_format = '0.00"%"'
                cell.font = Font(name="Arial", size=9, bold=True,
                                 color="375623" if val >= 0 else "9C0006")

            elif key in ("entry_price", "exit_price", "stop_loss", "target",
                         "risk_rs", "reward_rs"):
                cell.number_format = '#,##0.00'

            elif key == "direction":
                cell.font = Font(name="Arial", size=9, bold=True,
                                 color="375623" if val == "BUY" else "9C0006")

    # Column widths (26 columns)
    widths = [22, 12, 10, 10,  7, 12,  9,
              10, 10, 10, 10,  6,
              13, 12,            # Gross P&L ₹ + %
              12,                # Brokerage
              13, 12,            # Net P&L ₹ + %
              10, 10,  6,
              20,  7, 14, 16, 11, 25]
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"


def _build_daily_summary_sheet(ws, rows: dict):
    """Aggregate trades by day — shows Gross, Brokerage, Net P&L."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from collections import defaultdict

    hdr_fill = PatternFill("solid", start_color="2E75B6")
    hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    thin     = Side(style="thin", color="BFBFBF")
    border   = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = [
        "Date", "Total Trades", "Wins", "Losses", "Break Even", "Win Rate",
        "Gross P&L (₹)", "Brokerage (₹)", "Net P&L (₹)",
        "Best Trade Net (₹)", "Worst Trade Net (₹)", "Avg Net P&L (₹)",
        "Open Positions"
    ]

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border    = border
    ws.row_dimensions[1].height = 30

    by_date = defaultdict(list)
    for trade in rows.values():
        by_date[trade["date"]].append(trade)

    BASE_FONT = Font(name="Arial", size=9)
    ALT_FILL  = PatternFill("solid", start_color="EBF3FB")

    for row_idx, date_str in enumerate(sorted(by_date.keys()), start=2):
        trades = by_date[date_str]
        closed = [t for t in trades if t["status"] != "OPEN"]
        wins   = [t for t in closed if t["status"] == "WIN"]
        losses = [t for t in closed if t["status"] == "LOSS"]
        be     = [t for t in closed if t["status"] == "BREAKEVEN"]
        open_  = [t for t in trades if t["status"] == "OPEN"]

        gross_pnls = [t["pnl_rs"] for t in closed
                      if isinstance(t.get("pnl_rs"), (int, float))]
        brok_vals  = [t["brokerage_rs"] for t in closed
                      if isinstance(t.get("brokerage_rs"), (int, float))]
        net_pnls   = [_effective_pnl(t) for t in closed]

        gross_total = sum(gross_pnls)
        brok_total  = sum(brok_vals)
        net_total   = sum(net_pnls)
        best        = max(net_pnls) if net_pnls else 0
        worst       = min(net_pnls) if net_pnls else 0
        avg         = net_total / len(net_pnls) if net_pnls else 0
        wr          = len(wins) / len(closed) * 100 if closed else 0

        row_fill = ALT_FILL if row_idx % 2 == 0 else None
        values = [
            date_str, len(trades), len(wins), len(losses), len(be),
            round(wr, 1),
            round(gross_total, 2), round(brok_total, 2), round(net_total, 2),
            round(best, 2), round(worst, 2), round(avg, 2),
            len(open_)
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font      = BASE_FONT
            cell.border    = border
            cell.alignment = Alignment(vertical="center")
            if row_fill:
                cell.fill = row_fill

            if col_idx == 7:   # Gross P&L — secondary, grey
                cell.number_format = '#,##0.00'
                cell.font = Font(name="Arial", size=9, color="595959")

            elif col_idx == 8:  # Brokerage — orange
                cell.number_format = '#,##0.00'
                cell.font = Font(name="Arial", size=9, color="9C5700")

            elif col_idx == 9:  # Net P&L — primary, bold coloured
                cell.number_format = '₹#,##0.00;[Red](₹#,##0.00)'
                cell.font = Font(name="Arial", size=9, bold=True,
                                 color="375623" if val >= 0 else "9C0006")

            elif col_idx in (10, 11, 12):
                cell.number_format = '#,##0.00'
            elif col_idx == 6:
                cell.number_format = '0.0"%"'

    # Totals row
    last = len(by_date) + 2
    ws.cell(row=last, column=1, value="TOTAL").font = Font(name="Arial", bold=True, size=10)
    for c in (2, 3, 4):
        ws.cell(row=last, column=c,
                value=f'=SUM({get_column_letter(c)}2:{get_column_letter(c)}{last-1})'
               ).font = Font(name="Arial", bold=True)
    # Net P&L total (col 9)
    net_cell = ws.cell(row=last, column=9,
                       value=f'=SUM(I2:I{last-1})')
    net_cell.number_format = '₹#,##0.00;[Red](₹#,##0.00)'
    net_cell.font = Font(name="Arial", bold=True, size=11)

    widths = [14, 13, 7, 8, 12, 10, 16, 14, 16, 18, 19, 16, 15]
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w
    ws.freeze_panes = "A2"


def _build_phase_performance_sheet(ws, rows: dict):
    """Aggregate trades by phase - uses Net P&L as performance metric."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from collections import defaultdict

    hdr_fill = PatternFill("solid", start_color="375623")
    hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    thin     = Side(style="thin", color="BFBFBF")
    border   = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = [
        "Phase", "Total Trades", "Wins", "Losses", "Win Rate",
        "Total Net P&L (Rs)", "Avg Net P&L/Trade (Rs)",
        "Best Trade (Rs)", "Worst Trade (Rs)",
        "Total Brokerage (Rs)", "Avg Score", "Open"
    ]

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = border
    ws.row_dimensions[1].height = 30

    by_phase = defaultdict(list)
    for trade in rows.values():
        by_phase[trade["phase"]].append(trade)

    BASE_FONT = Font(name="Arial", size=9)
    PHASE_COLORS = {
        "PH3": "D9EAD3", "PH5": "CFE2F3", "PH5A": "EAD1DC",
        "PH5_INTRA": "CFE2F3",
        "PH6": "FFF2CC", "PH8": "FCE5CD", "MCX": "D0E4F1",
    }

    for row_idx, phase in enumerate(sorted(by_phase.keys()), start=2):
        trades = by_phase[phase]
        closed = [t for t in trades if t["status"] != "OPEN"]
        wins   = [t for t in closed if t["status"] == "WIN"]
        open_  = [t for t in trades if t["status"] == "OPEN"]

        net_pnls  = [_effective_pnl(t) for t in closed]
        brok_vals = [t["brokerage_rs"] for t in closed
                     if isinstance(t.get("brokerage_rs"), (int, float))]
        scores    = [t["score"] for t in trades
                     if isinstance(t.get("score"), (int, float)) and t["score"] > 0]

        total_net  = sum(net_pnls)
        total_brok = sum(brok_vals)
        wr         = len(wins) / len(closed) * 100 if closed else 0
        avg_net    = total_net / len(net_pnls) if net_pnls else 0
        best       = max(net_pnls) if net_pnls else 0
        worst      = min(net_pnls) if net_pnls else 0
        avg_score  = sum(scores) / len(scores) if scores else 0

        fill_color = PHASE_COLORS.get(phase, "F2F2F2")
        row_fill   = PatternFill("solid", start_color=fill_color)

        values = [
            phase, len(trades), len(wins), len(closed) - len(wins), round(wr, 1),
            round(total_net, 2), round(avg_net, 2),
            round(best, 2), round(worst, 2),
            round(total_brok, 2), round(avg_score, 1), len(open_)
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font      = BASE_FONT
            cell.fill      = row_fill
            cell.border    = border
            cell.alignment = Alignment(vertical="center")
            if col_idx in (6, 7, 8, 9):
                cell.number_format = "#,##0.00"
                if col_idx == 6:
                    cell.font = Font(name="Arial", size=9, bold=True,
                                     color="375623" if val >= 0 else "9C0006")
            elif col_idx == 10:
                cell.number_format = "#,##0.00"
                cell.font = Font(name="Arial", size=9, color="9C5700")
            elif col_idx == 5:
                cell.number_format = "0.0" + chr(34) + "%" + chr(34)

    widths = [12, 13, 7, 8, 10, 18, 20, 16, 17, 18, 11, 7]
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w
    ws.freeze_panes = "A2"


def _build_month_overview_sheet(ws, rows: dict):
    """High-level month overview - v2.0.0 shows gross + net P&L."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    TITLE_FONT = Font(name="Arial", bold=True, size=14, color="1F4E79")
    LABEL_FONT = Font(name="Arial", bold=True, size=10)
    VALUE_FONT = Font(name="Arial", size=10)
    HDR_FILL   = PatternFill("solid", start_color="BDD7EE")
    thin       = Side(style="thin", color="BFBFBF")
    border     = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 22

    ws["A1"] = "PAPER TRADE - 1 MONTH VERIFICATION (Net P&L v2.0)"
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = PatternFill("solid", start_color="DEEAF1")

    ws["A2"] = "Period Start: 2026-04-20  |  P&L shown NET of brokerage"
    ws["A2"].font = Font(name="Arial", italic=True, size=9, color="595959")

    all_closed = [t for t in rows.values() if t["status"] != "OPEN"]
    all_open   = [t for t in rows.values() if t["status"] == "OPEN"]
    wins       = [t for t in all_closed if t["status"] == "WIN"]
    losses     = [t for t in all_closed if t["status"] == "LOSS"]

    gross_pnls = [t["pnl_rs"] for t in all_closed
                  if isinstance(t.get("pnl_rs"), (int, float))]
    brok_vals  = [t["brokerage_rs"] for t in all_closed
                  if isinstance(t.get("brokerage_rs"), (int, float))]
    net_pnls   = [_effective_pnl(t) for t in all_closed]

    gross_total = sum(gross_pnls)
    brok_total  = sum(brok_vals)
    net_total   = sum(net_pnls)
    wr          = len(wins) / len(all_closed) * 100 if all_closed else 0
    avg_net     = net_total / len(net_pnls) if net_pnls else 0
    best_net    = max(net_pnls) if net_pnls else 0
    worst_net   = min(net_pnls) if net_pnls else 0
    dates       = sorted(set(t["date"] for t in rows.values()))

    sign = "+" if net_total >= 0 else ""
    metrics = [
        ("", ""),
        ("METRIC", "VALUE"),
        ("Total Trades", len(rows)),
        ("Closed Trades", len(all_closed)),
        ("Open Positions", len(all_open)),
        ("Wins (net basis)", len(wins)),
        ("Losses (net basis)", len(losses)),
        ("Win Rate", f"{wr:.1f}%"),
        ("", ""),
        ("Gross P&L", f"Rs {gross_total:,.2f}"),
        ("Total Brokerage Paid", f"Rs {brok_total:,.2f}"),
        ("Net P&L (after brokerage)", f"Rs {sign}{net_total:,.2f}"),
        ("Avg Net P&L / Trade", f"Rs {avg_net:,.2f}"),
        ("Best Trade (Net)", f"Rs {best_net:,.2f}"),
        ("Worst Trade (Net)", f"Rs {worst_net:,.2f}"),
        ("", ""),
        ("Trading Days", len(dates)),
        ("First Trade Date", dates[0] if dates else "-"),
        ("Last Trade Date",  dates[-1] if dates else "-"),
    ]

    for row_idx, (label, value) in enumerate(metrics, start=3):
        c_label = ws.cell(row=row_idx, column=1, value=label)
        c_value = ws.cell(row=row_idx, column=2, value=value)

        if label == "METRIC":
            c_label.font = Font(name="Arial", bold=True, size=10, color="1F4E79")
            c_value.font = Font(name="Arial", bold=True, size=10, color="1F4E79")
            c_label.fill = HDR_FILL
            c_value.fill = HDR_FILL
        elif label == "Net P&L (after brokerage)":
            c_label.font = LABEL_FONT
            c_value.font = Font(name="Arial", size=12, bold=True,
                                color="375623" if net_total >= 0 else "9C0006")
        elif label == "Total Brokerage Paid":
            c_label.font = LABEL_FONT
            c_value.font = Font(name="Arial", size=10, color="9C5700")
        elif label == "Gross P&L":
            c_label.font = LABEL_FONT
            c_value.font = Font(name="Arial", size=10, color="595959")
        elif label:
            c_label.font = LABEL_FONT
            c_value.font = VALUE_FONT

        if label and label != "METRIC":
            c_label.border = border
            c_value.border = border
